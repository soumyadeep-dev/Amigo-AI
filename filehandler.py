import pandas as pd
from pptx import Presentation
from PyPDF2 import PdfReader
import io
import re
from datetime import date
from database import insert_client_from_file, insert_log_for_client

# Keep your existing column variations - they are excellent for UX
NAME_COLS    = ["name", "client", "client name", "company", "business"]
PROJECT_COLS = ["project", "project name", "work", "service", "job"]
STATUS_COLS  = ["status", "project status", "state"]
PAYMENT_COLS = ["payment", "payment status", "paid", "payment_status"]
CONTACT_COLS = ["last contact", "last_contact", "date", "contact date", "last seen"]
NOTES_COLS   = ["notes", "note", "update", "updates", "comments", "description"]
INVOICE_COLS = ["invoice", "invoice amount", "amount", "fee", "price", "cost", "invoice_amount"]

def clean_extracted_text(text):
    """Removes excessive newlines and tabs to help the Vector model process better."""
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def match_col(columns, options):
    columns_lower = [c.lower().strip() for c in columns]
    for option in options:
        if option in columns_lower:
            return columns[columns_lower.index(option)]
    return None

def try_insert_as_clients(df):
    cols = df.columns.tolist()
    name_col    = match_col(cols, NAME_COLS)
    project_col = match_col(cols, PROJECT_COLS)

    if not name_col or not project_col:
        return False, 0

    status_col  = match_col(cols, STATUS_COLS)
    payment_col = match_col(cols, PAYMENT_COLS)
    contact_col = match_col(cols, CONTACT_COLS)
    notes_col   = match_col(cols, NOTES_COLS)
    invoice_col = match_col(cols, INVOICE_COLS)

    inserted = 0
    today    = str(date.today())

    for _, row in df.iterrows():
        name    = str(row[name_col]).strip()
        project = str(row[project_col]).strip()
        if not name or name == "nan": continue

        status         = str(row[status_col]).strip()  if status_col  and str(row[status_col])  != "nan" else "in progress"
        payment_status = str(row[payment_col]).strip() if payment_col and str(row[payment_col]) != "nan" else "pending"
        last_contact   = str(row[contact_col]).strip() if contact_col and str(row[contact_col]) != "nan" else today

        invoice = 0
        if invoice_col and str(row[invoice_col]) != "nan":
            try:
                invoice = float(str(row[invoice_col]).replace(',', '').replace('₹', '').replace('Rs', '').replace('rs', '').strip())
            except: invoice = 0

        insert_client_from_file(name, project, status, last_contact, payment_status, invoice)

        if notes_col and str(row[notes_col]) != "nan":
            note = str(row[notes_col]).strip()
            insert_log_for_client(name, note, today)
        inserted += 1

    return True, inserted

def parse_file(uploaded_file):
    filename = uploaded_file.name
    filetype = filename.split(".")[-1].lower()
    content  = ""
    inserted = 0
    was_client_data = False

    try:
        if filetype in ["xlsx", "xls", "csv"]:
            df = pd.read_csv(uploaded_file) if filetype == "csv" else pd.read_excel(uploaded_file)
            was_client_data, inserted = try_insert_as_clients(df)
            # For tabular data, we provide a clean summary for the Brain
            content = f"Spreadsheet {filename} contains {len(df)} rows. Key Columns: {', '.join(df.columns.tolist())}"

        elif filetype == "pptx":
            prs = Presentation(uploaded_file)
            text_runs = []
            for i, slide in enumerate(prs.slides):
                slide_text = " ".join([shape.text.strip() for shape in slide.shapes if hasattr(shape, "text")])
                if slide_text:
                    text_runs.append(f"[Slide {i+1}]: {slide_text}")
            content = clean_extracted_text("\n".join(text_runs))

        elif filetype == "pdf":
            reader = PdfReader(uploaded_file)
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages_text.append(f"[Page {i+1}]: {text}")
            content = clean_extracted_text("\n".join(pages_text))

        elif filetype == "txt":
            content = clean_extracted_text(uploaded_file.read().decode("utf-8"))

        else:
            content = f"Unsupported file type: {filetype}"

    except Exception as e:
        content = f"Error parsing file: {str(e)}"

    return filetype, content, was_client_data, inserted